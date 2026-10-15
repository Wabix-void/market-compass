"""演示 PPT 版式自检：越界、形状重叠、文本溢出估算。

用法：
    python scripts/check_deck_layout.py [docs/参赛材料/02-演示PPT.pptx]

说明：这是**几何估算**，不能替代人眼看渲染图，但能稳定抓住三类常见事故：
1. 形状超出画布；
2. 两个形状部分重叠（如说明文字压在卡片上）——"完全包含"关系视为有意设计，不报；
3. 文本框容量不足（按段落字号加权估算所需高度）。
"""

from __future__ import annotations

import math
import sys
from pathlib import Path

from pptx import Presentation

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT = PROJECT_ROOT / "docs" / "参赛材料" / "02-演示PPT.pptx"
EMU_IN = 914400.0
OVERLAP_RATIO = 0.08      # 交集占较小形状面积的比例超过此值即报告
OVERLAP_MIN_AREA = 0.02   # 绝对面积下限（平方英寸），过滤像素级接触


def _rect(shape) -> tuple[float, float, float, float] | None:
    if shape.left is None or shape.top is None or shape.width is None or shape.height is None:
        return None
    left, top = shape.left / EMU_IN, shape.top / EMU_IN
    return left, top, left + shape.width / EMU_IN, top + shape.height / EMU_IN


def _area(box) -> float:
    return max(box[2] - box[0], 0) * max(box[3] - box[1], 0)


def _intersection(a, b) -> float:
    width = min(a[2], b[2]) - max(a[0], b[0])
    height = min(a[3], b[3]) - max(a[1], b[1])
    return width * height if width > 0 and height > 0 else 0.0


def _is_contained(inner, outer) -> bool:
    return (inner[0] >= outer[0] - 0.01 and inner[1] >= outer[1] - 0.01
            and inner[2] <= outer[2] + 0.01 and inner[3] <= outer[3] + 0.01)


def _text_need_in(shape) -> float:
    width_in = shape.width / EMU_IN
    need_pt = 0.0
    for para in shape.text_frame.paragraphs:
        text = "".join(run.text for run in para.runs)
        size = max((r.font.size.pt for r in para.runs if r.font.size), default=12.0)
        space = para.space_after.pt if para.space_after else 0.0
        if not text:
            need_pt += size * 1.2 + space
            continue
        # 中文按整字宽、西文按 0.55 倍宽估算
        width_pt = sum((size if ord(ch) > 127 else size * 0.55) for ch in text)
        lines = max(1, math.ceil(width_pt / max(width_in * 72, 1)))
        need_pt += lines * size * 1.35 + space
    return need_pt / 72


def check(path: Path) -> list[str]:
    prs = Presentation(str(path))
    page_area = _area((0, 0, prs.slide_width / EMU_IN, prs.slide_height / EMU_IN))
    issues: list[str] = []

    for index, slide in enumerate(prs.slides, start=1):
        boxes = []
        for shape in slide.shapes:
            box = _rect(shape)
            if box is None:
                continue
            label = (shape.text_frame.text[:24].replace("\n", " ")
                     if shape.has_text_frame and shape.text_frame.text.strip()
                     else f"<{shape.shape_type}>")
            boxes.append((box, label, shape))

            if box[0] < -0.01 or box[1] < -0.01 \
                    or box[2] > prs.slide_width / EMU_IN + 0.01 \
                    or box[3] > prs.slide_height / EMU_IN + 0.01:
                issues.append(f"P{index} 越界: {label!r} {box}")

            if shape.has_text_frame and shape.text_frame.text.strip():
                height_in = shape.height / EMU_IN
                need = _text_need_in(shape)
                if need > height_in + 0.1:
                    issues.append(f"P{index} 文本容量不足: 需 {need:.2f}in / 有 {height_in:.2f}in"
                                  f" -> {label!r}")

        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                box_a, label_a, _ = boxes[i]
                box_b, label_b, _ = boxes[j]
                if _area(box_a) > page_area * 0.9 or _area(box_b) > page_area * 0.9:
                    continue                      # 整页背景，跳过
                if _is_contained(box_a, box_b) or _is_contained(box_b, box_a):
                    continue                      # 完全包含 = 有意设计（卡片内的色条/文字）
                overlap = _intersection(box_a, box_b)
                if overlap < OVERLAP_MIN_AREA:
                    continue
                ratio = overlap / max(min(_area(box_a), _area(box_b)), 1e-6)
                if ratio >= OVERLAP_RATIO:
                    issues.append(f"P{index} 形状重叠 {ratio * 100:.0f}%: "
                                  f"{label_a!r} × {label_b!r}")
    return issues


def main() -> int:
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT
    if not path.exists():
        print(f"找不到文件：{path}")
        return 1
    prs = Presentation(str(path))
    issues = check(path)
    print(f"文件：{path.name}    页数：{len(prs.slides)}")
    print(f"版式自检：发现 {len(issues)} 处问题")
    for item in issues:
        print("  -", item)
    if not issues:
        print("  ✓ 无越界、无形状重叠、无明显文本溢出")
    return 0 if not issues else 2


if __name__ == "__main__":
    sys.exit(main())
